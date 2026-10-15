"""Build the takeoff / estimate workbook for the SFH Development spec home (Ouachita Parish, LA).

Usage:  python3 scripts/build_estimate.py
Writes: output/SFH_Ouachita_Takeoff_Estimate.xlsx (formula-driven; recalculated by LibreOffice
        afterwards by scripts/build_all.sh so cached values exist for viewers that don't calculate).

Quantities: measured from the 5/31/2023 progress drawings (A100, A101, A401, A701) and the
Schematic Design Notes; every row carries its source sheet, calculation, basis and confidence.
Pricing: estimator budget unit costs for the Monroe / Ouachita Parish market, price date
2026-10-15. They are NOT vendor or subcontractor quotes.
"""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule

sys.path.insert(0, os.path.dirname(__file__))
import model_data as md  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..", "output", "SFH_Ouachita_Takeoff_Estimate.xlsx")
PRICE_DATE = "2026-10-15"
SRC = "Estimator budget unit cost, NE Louisiana (Monroe/Ouachita) Q4-2026 - not a quote; verify with local supplier/sub"
MARKET = "Monroe / West Monroe / Ouachita Parish, LA"

# ------------------------------------------------------------------ derived geometry (documented)
W, ZH, ZP, CD, PL = md.W, md.Z_HOUSE, md.Z_PORCH, md.CARPORT_D, md.PLATE
SF = md.SLOPE_FACTOR
ET = md.EXT_T
PORCH_D = ZP - ZH
ROOF_L = ZP + CD + 2 * md.RAKE_OH                     # 96.08'
ROOF_RUN = W / 2 + md.EAVE_OH                          # 11.5'
GABLE = 0.5 * W * (W / 2 * md.PITCH)                   # 55.13 SF each
PART_LEN = sum(math.dist((a, b), (c, d)) for a, b, c, d in md.PARTITIONS.values())
EXT_LEN = 2 * ZP + 2 * W
EXT_IN_PERIM = 2 * (ZH - 2 * ET) + 2 * (W - 2 * ET)   # interior face of conditioned exterior walls
WIN_A = sum(o[4] * o[5] for o in md.WINDOWS)          # 221
EXT_DOOR_A = sum(d[4] * d[5] for d in md.DOORS if d[6].startswith("ext"))  # 72
R3_A = 2.667 * 6.667
R4_A = 6.0 * 6.667
CEIL_ROOMS = [(n, (x1 - x0) * (z1 - z0), ch) for n, x0, z0, x1, z1, ch in md.ROOMS if ch]
CEIL_AREA = sum(a for _, a, _ in CEIL_ROOMS)
CEIL_8 = sum(a for _, a, ch in CEIL_ROOMS if ch == 8.0)
STOR = [r for r in md.ROOMS if r[0] == "STORAGE"][0]
STOR_PAINT_EX = 2 * ((STOR[3] - STOR[1]) + (STOR[4] - STOR[2])) * 8 + (STOR[3] - STOR[1]) * (STOR[4] - STOR[2])

r2 = lambda v: round(v, 2)

# ------------------------------------------------------------------ takeoff + cost rows
# fields: id, code, desc, spec, sheet, loc, ttype, count, L, Wd, Dp, factor, deduct, unit, waste, rnd,
#         calc, basis, conf, notes, (mat, lab, sub), net_formula(optional)
ROWS = []


def T(id, desc, spec, sheet, loc, ttype, unit, cost, count=None, L=None, Wd=None, Dp=None, factor=None,
      deduct=None, waste=0.0, rnd="N", calc="", basis="Measured", conf="High", notes="", net=None, est_desc=None):
    ROWS.append(dict(id=id, code=id[:2], desc=desc, spec=spec, sheet=sheet, loc=loc, ttype=ttype, count=count,
                     L=L, Wd=Wd, Dp=Dp, factor=factor, deduct=deduct, unit=unit, waste=waste, rnd=rnd, calc=calc,
                     basis=basis, conf=conf, notes=notes, cost=cost, net=net, est_desc=est_desc or desc))


LS = "LS"
# ---- 01 General requirements
T("01-100", "Building permit & plan review", "Ouachita Parish / municipal residential permit", "—", "Project", "Allowance", LS, (0, 0, 2600), count=1, basis="Allowance", conf="Medium", notes="Jurisdiction depends on final address (City of Monroe, West Monroe or unincorporated Ouachita). RFI-02.")
T("01-105", "Engineering: foundation design, carport/porch beams, truss review", "Licensed LA engineer letter/stamp", "—", "Project", "Allowance", LS, (0, 0, 1800), count=1, basis="Allowance", conf="Medium", notes="No structural drawings in set. RFI-05/06/07.")
T("01-106", "Geotechnical investigation (borings, PI, slab recommendations)", "Recommended for NE LA expansive clays", "—", "Site", "Allowance", LS, (0, 0, 1900), count=1, basis="Allowance", conf="Medium", notes="RFI-04")
T("01-110", "Survey: boundary, staking, form-board, final as-built", "", "A100", "Site", "Allowance", LS, (0, 0, 1850), count=1, basis="Allowance", conf="Medium")
T("01-120", "Temporary power pole, connection & usage", "", "—", "Site", "Allowance", LS, (250, 0, 900), count=1, basis="Allowance", conf="Medium")
T("01-125", "Temporary water & usage", "", "—", "Site", "Allowance", LS, (0, 0, 300), count=1, basis="Allowance", conf="Medium")
T("01-130", "Portable toilet", "", "—", "Site", "Count", "MO", (0, 0, 165), count=5, basis="Assumed", conf="Medium", calc="5 month schedule", notes="Schedule ASSUMED 5 months.")
T("01-140", "Dumpster, 20 CY, haul & dump", "", "—", "Site", "Count", "EA", (0, 0, 485), count=5, basis="Assumed", conf="Medium", calc="~1 pull per month")
T("01-150", "Builder's risk insurance", "", "—", "Project", "Allowance", LS, (0, 0, 1250), count=1, basis="Allowance", conf="Medium")
T("01-160", "Site supervision (part-time superintendent)", "", "—", "Project", "Count", "WK", (0, 650, 0), count=20, basis="Assumed", conf="Medium", calc="20 weeks x part-time super")
T("01-170", "Final construction clean incl. windows", "", "A101", "All areas", "Area", "SF", (0, 0, 0.35), count=1, L=W, Wd=ZP + CD, calc="21' x 94.08' under roof", basis="Calculated")
T("01-175", "Periodic jobsite cleanup", "", "—", "Project", "Allowance", LS, (0, 900, 0), count=1, basis="Allowance", conf="Medium")
T("01-180", "Finished-slab protection (Ram Board/taped) - smooth trowel floors", "Design notes: concrete floors are the finish", "Notes", "Conditioned", "Area", "SF", (0.38, 0.08, 0), L=W, Wd=ZH, waste=0.10, calc="21' x 64.17'", basis="Calculated", notes="Required because finished floor is exposed concrete.")
T("01-190", "Termite pre-treatment & LDAF certificate", "Louisiana pre-construction termite treatment", "—", "Slab", "Allowance", LS, (0, 0, 525), count=1, basis="Allowance", conf="Medium")
T("01-200", "Small tools & consumables", "", "—", "Project", "Allowance", LS, (600, 0, 0), count=1, basis="Allowance", conf="Medium")
T("01-210", "Energy code testing (blower door / duct leakage)", "", "—", "Project", "Allowance", LS, (0, 0, 450), count=1, basis="Allowance", conf="Medium", notes="Verify requirement with AHJ (RFI-03).")

# ---- 31 Earthwork
T("31-100", "Clear & grub, strip topsoil (building & drive area)", "", "A100", "Site", "Allowance", LS, (0, 0, 2800), count=1, basis="Allowance", conf="Low", notes="Lot condition unknown (address TBD). RFI-01.")
T("31-120", "Select fill building pad, 12\" avg, compacted (loose CY incl. 25% shrink)", "Slab footprint + 5' each side", "A101", "Building pad", "Volume", "CY", (18, 0, 7.5), L=W + 10, Wd=ZP + CD + 10, Dp=1.0, factor=1 / 27, waste=0.25, rnd="Y", calc="(21+10) x (94.08+10) x 1.0' / 27", basis="Assumed", conf="Low", notes="Pad height ASSUMED 12\" avg; set by survey/geotech & flood zone. RFI-04.")
T("31-130", "Fine grade building pad", "", "A101", "Building pad", "Area", "SF", (0, 0, 0.18), L=W + 10, Wd=ZP + CD + 10, calc="31' x 104.08'", basis="Calculated")
T("31-140", "Final grading & topsoil spread", "", "A100", "Lot", "Allowance", LS, (0, 0, 1800), count=1, basis="Allowance", conf="Medium")
T("31-150", "Silt fence / erosion control", "", "A100", "Lot perimeter (front & sides)", "Length", "LF", (0.9, 1.6, 0), L=50 + 2 * 90, calc="front 50' + 2 sides x 90'", basis="Assumed", conf="Medium")

# ---- 03 Concrete
T("03-100", "Slab-on-grade 4\" (house, porch, carport) - concrete", "3000 psi ready-mix", "A101", "Footprint 21' x 94.08'", "Volume", "CY", (168, 0, 0), L=W, Wd=ZP + CD, Dp=md.SLAB_T, factor=1 / 27, waste=0.05, calc="21' x (70.08'+24') x 0.333' / 27", basis="Calculated", conf="High", notes="Slab thickness ASSUMED 4\" (no structural). Footprint DIM.")
T("03-110", "Perimeter grade beam 12\" w x 20\" below slab - concrete", "Monolithic", "A101", "Perimeter", "Volume", "CY", (168, 0, 0), L=2 * (W + ZP + CD), Wd=1.0, Dp=20 / 12, factor=1 / 27, waste=0.08, calc="2 x (21 + 94.08) LF x 1.0' x 1.667' / 27", basis="Assumed", conf="Medium", notes="Beam size ASSUMED typical NE LA; engineer to confirm.")
T("03-120", "Interior grade beams 12\" w x 16\" below slab - concrete", "", "A101", "5 transverse + 1 longitudinal", "Volume", "CY", (168, 0, 0), L=5 * (W - 2) + (ZP + CD - 2), Wd=1.0, Dp=16 / 12, factor=1 / 27, waste=0.08, calc="(5 x 19' + 92.08') x 1.0' x 1.333' / 27", basis="Assumed", conf="Low", notes="Layout ASSUMED ~20' grid.")
T("03-130", "Vapor retarder 10 mil poly, lapped & taped", "", "A101", "Under slab", "Area", "SF", (0.11, 0.06, 0), L=W, Wd=ZP + CD, waste=0.15, calc="21' x 94.08' (+15% laps/beam sides)", basis="Calculated")
T("03-140", "Rebar #5 continuous in grade beams (4 bars)", "Grade 60", "A101", "All beams", "Length", "LF", (0.85, 0.35, 0), count=4, L=2 * (W + ZP + CD) + 5 * (W - 2) + (ZP + CD - 2), waste=0.10, calc="4 bars x (230.17 + 187.08) LF", basis="Assumed", conf="Medium")
T("03-150", "Rebar #3 @ 18\" O.C. each way in slab", "Grade 60", "A101", "Slab", "Length", "LF", (0.32, 0.18, 0), L=W, Wd=ZP + CD, factor=2 / 1.5, waste=0.10, calc="1975.7 SF x (2 / 1.5 ft) LF/SF", basis="Assumed", conf="Medium")
T("03-160", "Forms, chairs, stakes, ties, edge/expansion material", "", "A101", "Slab", "Area", "SF", (0.22, 0, 0), L=W, Wd=ZP + CD, calc="slab area", basis="Calculated")
T("03-170", "Foundation labor: form, set steel, place & finish (sub)", "", "A101", "Slab", "Area", "SF", (0, 3.0, 0), L=W, Wd=ZP + CD, calc="slab area", basis="Calculated")
T("03-175", "Concrete pump truck", "", "—", "Slab pour", "Allowance", LS, (0, 0, 1150), count=1, basis="Allowance", conf="Medium")
T("03-180", "Finished concrete floors: smooth trowel premium, grind, densifier & sealer", "Design notes: Concrete, smooth trowel (all rooms)", "Notes / A101", "Conditioned", "Area", "SF", (0, 0, 2.4), L=W, Wd=ZH, calc="21' x 64.17'", basis="Calculated", conf="High")
T("03-185", "Porch slab: smooth trowel, slope away, sealer", "Design notes", "Notes / A101", "Front porch", "Area", "SF", (0.35, 0.4, 0), L=W, Wd=PORCH_D, calc="21' x 5.92'", basis="Calculated")
T("03-190", "Anchor bolts / sill anchors at exterior walls", "", "A101", "Exterior walls", "Count", "EA", (2.75, 1.5, 0), count=46, calc="182.2 LF / 6' O.C. + 2 per plate piece (~8 pcs)", basis="Calculated", conf="Medium")
T("03-200", "Saw-cut control joints", "", "A101", "Slab", "Length", "LF", (0, 0, 1.75), L=120, calc="~5 joints x 21' + porch", basis="Assumed", conf="Medium")
T("03-220", "Carport post footings 24\"x24\"x24\" w/ post base", "", "A101", "Carport corners", "Count", "EA", (55, 0, 240), count=2, calc="2 posts shown", basis="Measured", conf="Medium", notes="Footing size ASSUMED.")

# ---- 06 Rough carpentry (material) + framing labor
T("06-100", "PT 2x6 bottom/sill plate - exterior walls", "", "A101", "Exterior walls", "Length", "LF", (1.15, 0, 0), L=r2(EXT_LEN), waste=0.10, calc="2 x 70.08 + 2 x 21", basis="Calculated")
T("06-110", "2x6 double top plate - exterior walls", "", "A101", "Exterior walls", "Length", "LF", (0.95, 0, 0), count=2, L=r2(EXT_LEN), waste=0.10, calc="2 plates x 182.17 LF", basis="Calculated")
T("06-120", "2x6 x 104-5/8\" studs @ 24\" O.C. - exterior walls", "Design notes: 2x6 ext walls @ 24\" O.C.", "Notes / A101", "Exterior walls", "Count", "EA", (9.25, 0, 0), count=f"=ROUNDUP({EXT_LEN:.2f}/2,0)+4*3+10*2+14*4+28", waste=0.10, rnd="Y", calc="LF/2 + corners 4x3 + partition T's 10x2 + 14 openings x4 (king+jack) + ~28 cripples", basis="Calculated", conf="Medium")
T("06-130", "2x10 header stock, 2-ply - exterior openings", "", "A701", "Exterior openings", "Length", "LF", (1.75, 0, 0), count=2, L=r2(5 * 6.5 + 2 * 6.5 + 2.5 + 9.5 + 4.5 + 2 * 3.5), waste=0.10, calc="2 ply x Σ(opening width + 6\"); D windows ganged on one 9.5' header", basis="Calculated", conf="Medium", notes="Header sizes ASSUMED (no structural).")
T("06-140", "PT 2x4 bottom plate - interior partitions", "", "A101", "Partitions", "Length", "LF", (0.80, 0, 0), L=r2(PART_LEN), waste=0.10, calc="Σ partitions P1-P10 (SCALED/DIM)", basis="Measured", conf="Medium")
T("06-150", "2x4 double top plate - interior partitions", "", "A101", "Partitions", "Length", "LF", (0.62, 0, 0), count=2, L=r2(PART_LEN), waste=0.10, calc="2 x partition LF", basis="Calculated")
T("06-160", "2x4 x 104-5/8\" studs @ 24\" O.C. - partitions", "Design notes: 2x4 int walls @ 24\" O.C.", "Notes / A101", "Partitions", "Count", "EA", (5.75, 0, 0), count=f"=ROUNDUP({PART_LEN:.2f}/2,0)+10*2+12*2+7*4", waste=0.10, rnd="Y", calc="LF/2 + wall ends 10x2 + intersections 12x2 + 7 doors x4", basis="Calculated", conf="Medium")
T("06-170", "Interior door headers / misc 2x6", "", "A701", "Interior doors", "Length", "LF", (0.95, 0, 0), L=60, calc="7 doors + misc", basis="Assumed", conf="Medium")
T("06-180", "Blocking & backing (cabinets, tubs, accessories, TV, grab-bar ready)", "", "A101", "Interior", "Length", "LF", (0.95, 0, 0), L=250, basis="Allowance", conf="Medium")
T("06-190", "Roof trusses - common 6/12, 21' span, flat bottom chord @ 4' O.C.", "Design notes: cambered roof truss 4' O.C.", "Notes / A401", "Over conditioned area", "Count", "EA", (205, 0, 0), count=17, calc="64.17'/4 + 1 = 17", basis="Calculated", conf="Medium", notes="Truss profile/camber by manufacturer. RFI-06.")
T("06-191", "Roof trusses - vaulted (parallel-chord) 6/12 for sloped carport/porch ceilings", "A101 RCP: sloped ceilings 6\"/1'-0\"", "A101 RCP", "Carport & porch", "Count", "EA", (275, 0, 0), count=6, calc="carport 5 + porch 1", basis="Calculated", conf="Medium")
T("06-192", "End trusses at open gables (vaulted, w/ rake outlookers)", "", "A401", "Carport & porch ends", "Count", "EA", (300, 0, 0), count=2, calc="2 ends", basis="Calculated", conf="Medium")
T("06-200", "2x4 roof purlins @ 24\" O.C. (flat) over 48\" trusses", "Needed to span 5/8\" ZIP roof at 48\" O.C.", "Notes", "Roof", "Length", "LF", (0.62, 0, 0), count=14, L=r2(ROOF_L), waste=0.10, calc="2 slopes x 7 lines (12.86' slope / 2') x 96.08'", basis="Assumed", conf="Low", notes="RFI-06: confirm 48\" truss spacing vs sheathing span rating.")
T("06-210", "2x4 ceiling furring @ 24\" O.C. under trusses (conditioned)", "5/8\" GWB cannot span 48\"", "Notes", "Conditioned ceiling", "Length", "LF", (0.62, 0, 0), count=10, L=r2(ZH - 2 * ET), waste=0.10, calc="10 lines x 63.08'", basis="Assumed", conf="Medium")
T("06-215", "2x4 ceiling joists & ledgers for 8' ceilings", "A101 RCP: 8'-0\" ceilings", "A101 RCP", "Storage, closets, baths", "Length", "LF", (0.62, 0, 0), L=r2(CEIL_8 / 2 * 1.4), waste=0.10, calc=f"{CEIL_8:.0f} SF of 8' ceilings / 2' O.C. x 1.4 (ledgers)", basis="Calculated", conf="Medium")
T("06-220", "2x4 nailers for sloped carport & porch ceilings", "", "A101 RCP", "Carport & porch", "Length", "LF", (0.62, 0, 0), count=13, L=r2(CD + PORCH_D), waste=0.10, calc="13 lines x 29.92'", basis="Assumed", conf="Medium")
T("06-230", "Truss permanent lateral & diagonal bracing", "", "—", "Roof", "Length", "LF", (0.62, 0, 0), L=400, basis="Allowance", conf="Medium")
T("06-240", "2x6 sub-fascia", "", "A401", "Eaves & rakes", "Length", "LF", (0.95, 0, 0), L=r2(2 * ROOF_L + 4 * ROOF_RUN * SF), waste=0.10, calc="2 x 96.08 + 4 x 12.86", basis="Calculated")
T("06-245", "2x4 soffit framing / lookouts", "", "A401", "Eaves & rakes", "Length", "LF", (0.62, 0, 0), L=200, basis="Allowance", conf="Medium")
T("06-250", "LVL beams 3-1/2\" x 14\" (carport eaves 2 x 25', porch front 22')", "No structural; 24' clear span", "A101 / A401", "Carport & porch", "Length", "LF", (34, 0, 0), L=2 * 25 + 22, calc="2 x 25' + 22'", basis="Assumed", conf="Low", notes="Size ASSUMED - engineer to design (RFI-07).")
T("06-260", "6x6 PT posts w/ base & cap, wrapped", "", "A101", "Carport corners", "Count", "EA", (150, 0, 0), count=2, calc="2 posts shown on A101", basis="Measured", conf="Medium")
T("06-270", "ZIP System 7/16\" wall sheathing, 4x9 panels", "Design notes: Huberwood ZIP wall sheathing", "Notes / A401", "Exterior walls + gables", "Area", "SHT", (45, 0, 0), net=f"=({2 * ZP:.3f}*9+2*{W}*9+2*{GABLE:.2f})/36", waste=0.10, rnd="Y", calc="(2x70.08x9 + 2x21x9 + 2 gables 55.13) / 36 SF per sheet (openings not deducted - cut waste)", basis="Calculated")
T("06-280", "ZIP System 5/8\" roof sheathing, 4x8 panels", "Design notes: ZIP roof sheathing", "Notes / A401", "Roof", "Area", "SHT", (52, 0, 0), net=f"={ROOF_L:.3f}*{ROOF_RUN}*{SF:.4f}*2/32", waste=0.10, rnd="Y", calc="96.08' x 11.5' x 1.118 x 2 / 32", basis="Calculated")
T("06-290", "ZIP System tape 3.75\" x 90' rolls", "~7 panels per roll (Huber)", "Notes", "Wall & roof seams", "Count", "RL", (36, 0, 0), net="=({G:06-270}+{G:06-280})/7", waste=0.0, rnd="Y", calc="(wall + roof sheets) / 7", basis="Calculated")
T("06-295", "ZIP liquid flash / stretch tape at openings & penetrations", "", "A701", "Openings", "Allowance", LS, (350, 0, 0), count=1, basis="Allowance", conf="Medium")
T("06-300", "Structural connectors (truss ties, straps, hangers, post bases)", "", "—", "Structure", "Allowance", LS, (1150, 0, 0), count=1, basis="Allowance", conf="Medium")
T("06-310", "Nails, screws, adhesives, sill gasket", "", "—", "Structure", "Area", "SF", (0.45, 0, 0), L=W, Wd=ZP + CD, calc="area under roof", basis="Calculated")
T("06-320", "Framing labor - conditioned (walls, trusses, purlins, sheathing)", "", "A101", "Conditioned", "Area", "SF", (0, 7.25, 0), L=W, Wd=ZH, calc="21' x 64.17'", basis="Calculated")
T("06-325", "Framing labor - porch & carport (beams, posts, trusses, sheathing)", "", "A101", "Porch & carport", "Area", "SF", (0, 3.75, 0), L=W, Wd=CD + PORCH_D, calc="21' x 29.92'", basis="Calculated")
T("06-340", "Telehandler rental (truss & material set)", "", "—", "Project", "Allowance", LS, (0, 0, 1150), count=1, basis="Allowance", conf="Medium")

# ---- 07 Thermal & moisture
T("07-100", "Wall insulation R-21 batts (2x6)", "R-value ASSUMED", "A101", "Conditioned ext. walls", "Area", "SF", (0, 0, 1.15), net=f"={EXT_IN_PERIM:.2f}*9-({WIN_A}+{EXT_DOOR_A:.1f})", waste=0.05, calc="interior perimeter 166.0 LF x 9' - openings 293 SF", basis="Calculated", conf="High", notes="Verify R-values with energy code (RFI-03).")
T("07-110", "Ceiling insulation blown fiberglass R-38", "R-value ASSUMED", "A101 RCP", "Conditioned ceiling", "Area", "SF", (0, 0, 1.45), net=f"={CEIL_AREA:.2f}", calc="Σ room ceiling areas (A101 RCP)", basis="Calculated")
T("07-115", "Air sealing package (top plates, penetrations, rim)", "", "—", "Envelope", "Allowance", LS, (0, 0, 650), count=1, basis="Allowance", conf="Medium")
T("07-130", "Architectural asphalt shingles, black/charcoal", "Design notes", "Notes / A401", "Roof 6/12", "Area", "SQ", (118, 92, 0), net=f"={ROOF_L:.3f}*{ROOF_RUN}*{SF:.4f}*2/100", waste=0.10, calc="96.08 x 11.5 x 1.118 x 2 / 100", basis="Calculated", notes="ZIP roof serves as underlayment (Huber listing) - no felt included.")
T("07-140", "Hip & ridge cap shingles", "", "A401", "Ridge", "Length", "LF", (3.25, 1.25, 0), L=r2(ROOF_L), waste=0.05, calc="ridge length 96.08'", basis="Calculated")
T("07-145", "Ridge vent (over conditioned attic only)", "", "A401", "Ridge", "Length", "LF", (3.75, 1.5, 0), L=r2(ZH), calc="64.17'", basis="Assumed", conf="Medium")
T("07-150", "Fiber cement smooth panel siding 4x9, incl. Z-flash", "Design notes / A401 'Fiber cement smooth panel siding'", "A401", "W, E, N elevations + N gable", "Area", "SF", (2.15, 2.35, 0), net=f"=2*{ZP:.3f}*9+{W}*9+{GABLE:.2f}-(5*6*4+2*6*4+2*4+2*3*6.667)", waste=0.12, calc="2x70.08x9 + 21x9 + gable 55.13 - (A 120 + B 48 + C 8 + doors 09/10 40)", basis="Calculated")
T("07-155", "Panel joint trim / battens", "A401 joint lines", "A401", "Panel siding", "Length", "LF", (1.1, 1.25, 0), L=200, calc="~16 vertical joints x 9' + horizontals", basis="Assumed", conf="Medium")
T("07-160", "Fiber cement T&G siding - porch walls & front gable", "Notes: T&G (also says 'lap siding front porch')", "A401 Elev B / Notes", "Front porch", "Area", "SF", (2.85, 2.75, 0), net=f"={W}*9-(3*3*5+4*8)+2*{PORCH_D:.3f}*9+{GABLE:.2f}", waste=0.15, calc="21x9 - (D 45 + door 01 32) + porch sides 2x5.92x9 + gable 55.13", basis="Calculated", conf="Medium", notes="Lap vs T&G conflict - RFI-09.")
T("07-170", "Fiber cement T&G porch ceiling (sloped)", "Design notes", "A101 RCP / Notes", "Front porch", "Area", "SF", (3.1, 3.25, 0), L=W, Wd=PORCH_D, factor=SF, waste=0.15, calc="21 x 5.92 x 1.118 slope", basis="Calculated")
T("07-180", "Fiber cement panel ceiling - covered parking (sloped)", "Design notes", "A101 RCP / Notes", "Covered parking", "Area", "SF", (2.15, 2.6, 0), L=W, Wd=CD, factor=SF, waste=0.12, calc="21 x 24 x 1.118 slope", basis="Calculated")
T("07-190", "Fiber cement vented soffit, 12\"", "Overhang ASSUMED 12\"", "A401", "Eaves & rakes", "Length", "LF", (3.25, 3.0, 0), L=r2(2 * ROOF_L + 4 * ROOF_RUN * SF), waste=0.10, calc="2 x 96.08 + 4 x 12.86", basis="Assumed", conf="Medium")
T("07-200", "Fiber cement fascia 5/4x6, painted", "A401", "A401", "Eaves & rakes", "Length", "LF", (2.4, 2.1, 0), L=r2(2 * ROOF_L + 4 * ROOF_RUN * SF), waste=0.10, calc="2 x 96.08 + 4 x 12.86", basis="Calculated")
T("07-210", "Fiber cement trim 5/4x4 (corners & opening casings)", "", "A401", "Exterior", "Length", "LF", (1.75, 2.1, 0), L=r2(6 * 9 + 5 * 20 + 2 * 20 + 12 + 28 + 20 + 2 * (3 + 2 * 6.667)), waste=0.10, calc="6 corners x 9' + opening perimeters (A 100, B 40, C 12, D ganged 28, 01 20, 09/10 32.7)", basis="Calculated", conf="Medium")
T("07-220", "5\" seamless aluminum gutters", "Shown on A401", "A401", "Both eaves", "Length", "LF", (0, 0, 8.5), count=2, L=r2(ROOF_L), calc="2 x 96.08'", basis="Calculated")
T("07-225", "2x3 aluminum downspouts w/ splash blocks", "", "A401", "Corners", "Length", "LF", (0, 0, 8.0), count=4, L=10.5, calc="4 x 10.5'", basis="Assumed", conf="Medium")
T("07-230", "Drip edge", "", "A401", "Eaves & rakes", "Length", "LF", (0.85, 0.4, 0), L=r2(2 * ROOF_L + 4 * ROOF_RUN * SF), waste=0.10, calc="eaves + rakes", basis="Calculated")
T("07-235", "Starter strip", "", "A401", "Eaves", "Length", "LF", (0.75, 0.25, 0), count=2, L=r2(ROOF_L), waste=0.05, calc="2 x 96.08'", basis="Calculated")
T("07-240", "Pipe boots, roof flashing & sealant", "", "—", "Roof", "Allowance", LS, (225, 150, 0), count=1, basis="Allowance", conf="Medium")
T("07-250", "Exterior caulking & sealants", "", "—", "Exterior", "Allowance", LS, (450, 500, 0), count=1, basis="Allowance", conf="Medium")

# ---- 08 Openings
T("08-100", "Window Type A - vinyl sliding 6'-0\" x 4'-0\"", "A701", "A701 / A401 Elev D", "West wall (hall/kitchen/living)", "Count", "EA", (420, 95, 0), count=5, calc="schedule = plan = elevation (5)", basis="Measured")
T("08-110", "Window Type B - vinyl casement 6'-0\" x 4'-0\"", "A701", "A701 / A401 Elev C", "Master Bed, Bed #2", "Count", "EA", (925, 110, 0), count=2, basis="Measured", calc="schedule = plan (2)", notes="Egress window - verify clear opening (RFI-10).")
T("08-120", "Window Type C - vinyl single hung 2'-0\" x 4'-0\"", "A701", "A701 / A401 Elev A", "Laundry", "Count", "EA", (215, 85, 0), count=1, basis="Measured")
T("08-130", "Window Type D - vinyl single hung 3'-0\" x 5'-0\" (ganged x3, grids)", "A701", "A701 / A401 Elev B", "Living (porch wall)", "Count", "EA", (335, 95, 0), count=3, basis="Measured", notes="Tempered glazing likely near door 01 (RFI-11).")
T("08-200", "Door 01 (R1) exterior entry 4'-0\" x 8'-0\" w/ sidelite, prehung", "A701", "A701", "Living / porch", "Count", "EA", (1950, 275, 0), count=1, basis="Measured", conf="Medium", notes="Leaf vs sidelite split not dimensioned (RFI-12).")
T("08-210", "Door 09 (R2) exterior 3'-0\" x 6'-8\" half-lite, prehung", "A701", "A701", "Laundry from carport", "Count", "EA", (520, 185, 0), count=1, basis="Measured")
T("08-220", "Door 10 (R5) exterior 3'-0\" x 6'-8\" flush, prehung", "A701", "A701", "Storage from carport", "Count", "EA", (365, 185, 0), count=1, basis="Measured")
T("08-230", "Doors 02/04/05/06/08 (R3) interior 2'-8\" x 6'-8\" shaker, prehung", "A701", "A701", "Interior", "Count", "EA", (195, 95, 0), count=5, basis="Measured")
T("08-240", "Doors 03/07 (R4) interior bypass 6'-0\" x 6'-8\" 2-panel", "A701 / Notes", "A701", "Closets", "Count", "EA", (395, 145, 0), count=2, basis="Measured")
T("08-250", "Hardware - entry handleset & deadbolt (01)", "", "A701", "Door 01", "Count", "EA", (215, 45, 0), count=1, basis="Measured")
T("08-255", "Hardware - exterior keyed lever & deadbolt (09, 10)", "", "A701", "Doors 09/10", "Count", "EA", (95, 40, 0), count=2, basis="Measured")
T("08-260", "Hardware - privacy levers (02, 04, 05, 06)", "", "A701", "Bedrooms & baths", "Count", "EA", (32, 20, 0), count=4, basis="Assumed", conf="Medium")
T("08-265", "Hardware - passage lever (08)", "", "A701", "Laundry", "Count", "EA", (28, 20, 0), count=1, basis="Assumed", conf="Medium")
T("08-270", "Hardware - bypass pulls & track (03, 07)", "", "A701", "Closets", "Count", "EA", (22, 10, 0), count=2, basis="Measured")
T("08-275", "Door stops", "", "A701", "Swing doors", "Count", "EA", (5, 4, 0), count=8, basis="Calculated")

# ---- 09 Finishes
T("09-100", "1/2\" gypsum board - walls, hang & finish Level 3", "Design notes: level 3 (storage level 2)", "Notes / A101", "All walls", "Area", "SF", (0.58, 1.55, 0), net=f"={EXT_IN_PERIM:.2f}*9-({WIN_A}+{EXT_DOOR_A:.1f})+{PART_LEN:.2f}*9*2-(5*{R3_A:.3f}*2+2*{R4_A:.3f}*2)", waste=0.12, calc="ext. walls int. face 166.0x9 - 293 + partitions 138.6x9x2 - int. doors 338", basis="Calculated", notes="Walls taken to 9' in all rooms (8' ceiling rooms slightly over).")
T("09-105", "Moisture-resistant gypsum premium (baths & laundry)", "", "A101", "Baths, laundry", "Area", "SF", (0.18, 0, 0), L=955, calc="bath #2 275 + master bath 304 + laundry 376", basis="Calculated", conf="Medium")
T("09-110", "5/8\" gypsum board - ceilings, hang & finish Level 3", "", "A101 RCP", "Conditioned ceilings", "Area", "SF", (0.72, 1.75, 0), net=f"={CEIL_AREA:.2f}", waste=0.12, calc="Σ room ceiling areas", basis="Calculated")
T("09-120", "Drywall accessories (bead, mud, tape, screws)", "", "—", "Walls & ceilings", "Area", "SF", (0.12, 0, 0), net="={N:09-100}+{N:09-110}", calc="walls + ceilings", basis="Calculated")
T("09-200", "Baseboard 3-1/4\" MDF, painted", "Base not specified - ASSUMED", "A101", "All rooms", "Length", "LF", (1.1, 1.35, 0), net=f"={EXT_IN_PERIM:.2f}+2*{PART_LEN:.2f}-(5*2.667*2+2*6*2+10+2*5+20)", waste=0.10, calc="ext. 166 + partitions 2x138.6 - door openings 60.7 - tubs 10 - cabinets 20", basis="Assumed", conf="Medium")
T("09-210", "Door casing 2-1/4\" MDF, painted", "ASSUMED", "A701", "Interior doors", "Length", "LF", (0.85, 1.25, 0), net="=5*2*(2*6.833+2.667+0.5)+2*2*(2*6.833+6+0.5)+(2*8.17+4.5)+2*(2*6.83+3.5)", waste=0.10, calc="R3 168.3 + R4 80.7 + ext. doors interior 55.1", basis="Calculated", conf="Medium")
T("09-220", "Window stool & apron, MDF (drywall returns)", "ASSUMED", "A701", "Windows", "Length", "LF", (1.6, 2.5, 0), net="=2*(5*6.5+2*6.5+2.5+9.5)", waste=0.10, calc="2 x Σ(width + 6\"), D ganged", basis="Assumed", conf="Medium")
T("09-230", "Interior paint - walls & ceilings (primer + 2 coats)", "Storage not painted (level 2 only)", "Notes", "Interior", "Area", "SF", (0.28, 1.02, 0), net=f"={{N:09-100}}+{{N:09-110}}-{STOR_PAINT_EX:.1f}", calc="walls + ceilings - storage", basis="Calculated")
T("09-240", "Paint interior doors (both faces)", "", "A701", "Interior doors", "Count", "EA", (18, 65, 0), count=10, calc="5 R3 + 4 R4 panels + 1 (ext. door interior faces avg.)", basis="Calculated", conf="Medium")
T("09-250", "Paint interior trim", "", "—", "Trim", "Length", "LF", (0.1, 0.75, 0), net="={N:09-200}+{N:09-210}+{N:09-220}", calc="base + casing + stool/apron", basis="Calculated")
T("09-260", "Exterior paint - fiber cement siding, ceilings, soffit", "Primed FC field painted (ASSUMED)", "A401", "Exterior", "Area", "SF", (0.32, 0.95, 0), net="={N:07-150}+{N:07-160}+{N:07-170}+{N:07-180}+{N:07-190}", calc="panel + T&G + porch ceiling + carport ceiling + soffit", basis="Calculated", conf="Medium")
T("09-265", "Exterior paint - fascia & trim", "", "A401", "Exterior", "Length", "LF", (0.08, 0.65, 0), net="={N:07-200}+{N:07-210}+{N:07-155}", calc="fascia + trim + joint trim", basis="Calculated")
T("09-270", "Exterior doors paint/finish", "", "A701", "Doors 01/09/10", "Count", "EA", (25, 85, 0), count=3, basis="Measured")
T("09-280", "Interior caulk & touch-up", "", "—", "Interior", "Allowance", LS, (150, 250, 0), count=1, basis="Allowance", conf="Medium")

# ---- 10 Specialties
T("10-100", "Bath accessories (towel bar, ring, TP holder, hook)", "", "Notes", "Baths", "Count", "SET", (140, 85, 0), count=2, basis="Assumed", conf="Medium")
T("10-110", "Vanity mirrors", "Notes: mirror over vanity", "Notes", "Baths", "Count", "EA", (95, 45, 0), count=2, basis="Measured")
T("10-120", "Closet shelf & rod / laundry & storage shelving", "Notes: rod and shelf", "Notes / A101", "Closets, laundry, storage", "Length", "LF", (9, 7, 0), L=32, calc="master closet 10 + bed #2 closet 9.7 + laundry 6 + storage 6", basis="Measured", conf="Medium")
T("10-130", "Tub shower rods", "", "—", "Baths", "Count", "EA", (30, 20, 0), count=2, basis="Assumed", conf="Medium")
T("10-140", "House numbers & mailbox", "", "—", "Exterior", "Allowance", LS, (175, 75, 0), count=1, basis="Allowance", conf="Medium")

# ---- 11 Equipment
T("11-100", "Appliances (range, refrigerator, OTR microwave, dishwasher, washer/dryer)", "Design notes: supplied by owner", "Notes", "Kitchen, laundry", "Allowance", LS, (0, 0, 0), count=1, basis="Allowance", conf="High", notes="Owner furnished - $0 in contract.")
T("11-110", "Appliance installation & hookups (cords, DW, ice maker, OTR vent)", "", "Notes", "Kitchen, laundry", "Allowance", LS, (85, 350, 0), count=1, basis="Allowance", conf="Medium")

# ---- 12 Furnishings (casework)
T("12-100", "Base cabinets - frameless, shaker doors", "Design notes", "A101 / Notes", "Kitchen run + island", "Length", "LF", (285, 45, 0), net="=6.0+7.7-2.0", calc="back run 6.0 + island 7.7 - DW space 2.0 (SCALED)", basis="Measured", conf="Medium")
T("12-110", "Wall cabinets - frameless, shaker doors", "Design notes", "A101 / Notes", "Kitchen", "Length", "LF", (215, 40, 0), net="=2.9+3.1+3.4", calc="flanking range 2.9 + 3.1 + over refrigerator 3.4 (SCALED)", basis="Measured", conf="Medium")
T("12-120", "Tall pantry cabinet 30\" x 24\" x 84\"", "", "A101", "Kitchen", "Count", "EA", (875, 85, 0), count=1, basis="Measured", conf="Low", notes="Plan symbol at west end of run - verify (RFI-13).")
T("12-125", "Cabinet fillers, finished ends, island back panel", "", "A101", "Kitchen", "Allowance", LS, (350, 100, 0), count=1, basis="Allowance", conf="Medium")
T("12-130", "Bath vanities 30\" frameless shaker", "Design notes", "A101", "Bath #2, master bath", "Count", "EA", (425, 85, 0), count=2, basis="Measured", conf="Medium")
T("12-200", "Laminate countertops (post-form/custom)", "Design notes: laminate", "A101 / Notes", "Kitchen, island, vanities", "Area", "SF", (22, 9, 0), net="=6.0*2.125+7.9*3.0+2*2.5*1.83", waste=0.10, calc="back run 12.75 + island 23.7 + vanities 9.15", basis="Calculated", conf="Medium")
T("12-210", "Laminate 4\" backsplash", "", "A101", "Kitchen, vanities", "Length", "LF", (6, 3, 0), net="=6.0+2*2.1+2*(2.5+2*1.83)", calc="kitchen 10.2 + vanities 12.3", basis="Calculated", conf="Medium")

# ---- 22 Plumbing
T("22-100", "Plumbing rough-in (under-slab DWV, PEX supply, top-out)", "", "A101", "13 connections", "Count", "EA", (220, 480, 0), count=13, calc="kit. sink, DW, ice, 2 lav, 2 WC, 2 tub, washer, WH, 2 hose bibs", basis="Calculated", conf="Medium")
T("22-110", "Kitchen sink SS 33\" drop-in w/ chrome pull-down faucet", "Notes: chrome finish", "A101", "Island", "Count", "EA", (425, 0, 0), count=1, basis="Measured")
T("22-120", "Lavatory (drop-in) w/ chrome faucet & drain", "", "A101", "Baths", "Count", "EA", (195, 0, 0), count=2, basis="Measured")
T("22-130", "Water closet, elongated 2-pc w/ seat", "", "A101", "Baths", "Count", "EA", (285, 0, 0), count=2, basis="Measured")
T("22-140", "Bathtub 60\" w/ chrome tub/shower valve & trim", "Notes", "A101", "Baths", "Count", "EA", (520, 150, 0), count=2, basis="Measured")
T("22-145", "Tub surround, 3-pc acrylic", "Notes: bathtub surround", "Notes", "Baths", "Count", "EA", (375, 225, 0), count=2, basis="Measured")
T("22-150", "Water heater 50 gal electric w/ pan, exp. tank, T&P drain", "WH shown in storage (A101)", "A101", "Storage", "Count", "EA", (925, 350, 0), count=1, basis="Measured", conf="Medium", notes="Size/fuel ASSUMED.")
T("22-160", "Washer box, ice maker box, (2) frost-free hose bibs", "Notes: 2 hose bibs", "Notes", "Laundry, kitchen, exterior", "Allowance", LS, (190, 0, 0), count=1, basis="Measured")
T("22-190", "Fixture trim-out labor", "", "—", "Fixtures", "Count", "EA", (0, 145, 0), count=7, calc="sink, 2 lav, 2 WC, 2 tub", basis="Calculated")
T("22-195", "Plumbing permit & inspections", "", "—", "Project", "Allowance", LS, (0, 0, 275), count=1, basis="Allowance", conf="Medium")

# ---- 23 HVAC
T("23-100", "2-ton split heat pump w/ air handler & 5 kW heat strip", "Notes: 2-ton unit, air handler, ducted", "Notes", "Attic AH / exterior condenser", "Count", "EA", (4650, 0, 0), count=1, basis="Measured", conf="Medium", notes="Heat pump vs straight-cool + furnace ASSUMED (no gas shown). AH location RFI-14.")
T("23-110", "Duct system (plenum, R-8 flex, registers, returns)", "", "A101", "Conditioned", "Area", "SF", (1.35, 1.1, 0), L=W, Wd=ZH, calc="conditioned area", basis="Calculated")
T("23-120", "HVAC installation, line set, condensate, pad, thermostat, start-up", "", "—", "System", "Allowance", LS, (650, 2600, 0), count=1, basis="Allowance", conf="Medium")
T("23-130", "Bath exhaust fans, ducted to exterior", "Notes", "Notes / A101 RCP", "Baths", "Count", "EA", (95, 125, 0), count=2, basis="Measured")
T("23-140", "Dryer vent & wall cap", "", "A101", "Laundry", "Count", "EA", (55, 110, 0), count=1, basis="Assumed", conf="Medium")
T("23-150", "Microwave/range vent duct to exterior", "Notes: microwave over stove", "Notes", "Kitchen", "Count", "EA", (65, 125, 0), count=1, basis="Assumed", conf="Medium")
T("23-160", "Mechanical permit", "", "—", "Project", "Allowance", LS, (0, 0, 175), count=1, basis="Allowance", conf="Medium")

# ---- 26 Electrical
T("26-100", "200A service: meter/main, panel, grounding, surge protection", "", "—", "Storage / east wall", "Allowance", LS, (1750, 1250, 0), count=1, basis="Assumed", conf="Medium", notes="Service size ASSUMED (all-electric).")
T("26-110", "Branch rough-in wiring - conditioned (AFCI/GFCI, boxes, wire)", "", "A101", "Conditioned", "Area", "SF", (1.55, 2.1, 0), L=W, Wd=ZH, calc="conditioned area", basis="Calculated")
T("26-115", "Branch rough-in wiring - porch & carport", "", "A101", "Porch & carport", "Area", "SF", (0.45, 0.65, 0), L=W, Wd=CD + PORCH_D, calc="21 x 29.92", basis="Calculated")
T("26-120", "Devices: receptacles, switches, GFCI, plates", "", "A101", "All rooms", "Count", "EA", (7, 18, 0), count=64, calc="~38 receptacles + ~26 switches (code-minimum layout ASSUMED)", basis="Assumed", conf="Low", notes="No electrical plan (RFI-15).")
T("26-130", "Dedicated appliance/equipment circuits", "", "—", "Range, dryer, WH, HVAC(2), DW, MW, fridge", "Count", "EA", (95, 125, 0), count=8, basis="Calculated", conf="Medium")
T("26-140", "Smoke/CO combination detectors, hardwired, interconnected", "", "A101", "Bedrooms, hall, living", "Count", "EA", (48, 30, 0), count=4, basis="Assumed", conf="Medium")
T("26-300", "LED recessed 6\" wafer lights", "Notes / A101 RCP", "A101 RCP", "Per RCP", "Count", "EA", (24, 40, 0), count=sum(1 for l in md.LIGHTS if l[3] == "can"), calc="RCP symbol count (carport 4, laundry, storage, closet, hall 5, baths 4, closet, living 4, porch 2)", basis="Measured", conf="Medium")
T("26-310", "Pendant lights over island", "A101 RCP note", "A101 RCP", "Kitchen island", "Count", "EA", (120, 55, 0), count=2, basis="Measured")
T("26-320", "Ceiling fans w/ light kit & fan-rated box", "Notes", "A101 RCP", "Master, Bed #2, Living", "Count", "EA", (175, 85, 0), count=3, basis="Measured")
T("26-325", "Vanity lights", "Notes", "Notes", "Baths", "Count", "EA", (85, 45, 0), count=2, basis="Measured")
T("26-330", "Exterior lights (6 soffit downlights, front door sconce, rear entry light)", "Notes: down lighting from soffits", "Notes / A401", "Exterior", "Count", "EA", (55, 50, 0), count=8, basis="Assumed", conf="Medium")
T("26-340", "Low-voltage (data/TV 4 drops, doorbell)", "", "—", "Interior", "Allowance", LS, (225, 375, 0), count=1, basis="Allowance", conf="Medium")
T("26-350", "Electrical permit", "", "—", "Project", "Allowance", LS, (0, 0, 225), count=1, basis="Allowance", conf="Medium")

# ---- 32 Exterior improvements
DRIVE_L = ZP + CD + 25 + 10
T("32-100", "Concrete driveway 4\", 12' wide (lot + 10' ROW apron)", "A101: 12'-0\" drive", "A100 / A101", "West side", "Area", "SF", (3.35, 3.9, 0), L=12, Wd=r2(DRIVE_L), calc="12' x (24 + 70.08 + 25 front yard + 10 ROW)", basis="Measured", conf="Medium", notes="ROW length ASSUMED 10'; culvert see 32-130.")
T("32-110", "Concrete apron at covered parking (5' + 4' strips)", "A101", "A101", "West of carport", "Area", "SF", (3.35, 3.9, 0), L=5 + 4, Wd=CD, calc="9' x 24'", basis="Measured")
T("32-120", "Front walk / landing pad", "A101", "A101", "Front porch", "Area", "SF", (3.35, 3.9, 0), L=10.4, Wd=6.0, calc="10.4' x 6.0' (SCALED)", basis="Measured", conf="Medium")
T("32-130", "Driveway culvert & headwalls (if open-ditch street)", "", "A100", "Street", "Allowance", LS, (0, 0, 1200), count=1, basis="Allowance", conf="Low", notes="RFI-01")
T("32-200", "Shrubs, 3 gal (planting bed along west wall)", "A101 plant symbols", "A101", "West bed", "Count", "EA", (28, 15, 0), count=12, basis="Measured", conf="Medium")
T("32-205", "Bed prep, edging & mulch", "", "A101", "4' x 64.17' bed", "Area", "SF", (0.55, 0.45, 0), L=4, Wd=r2(ZH), calc="4' x 64.17'", basis="Measured")
T("32-207", "Shade tree, 30 gal", "A100/A101 tree symbol", "A101", "Front yard", "Count", "EA", (275, 125, 0), count=1, basis="Measured")
T("32-210", "Sod - front & east side yards", "", "A100", "Front/side", "Area", "SF", (0.55, 0.3, 0), L=1700, calc="front ~950 + east side 8' x 94' ~750", basis="Assumed", conf="Medium")
T("32-215", "Seed & straw - balance of disturbed lot", "", "A100", "Rear", "Area", "SF", (0.08, 0.06, 0), net=f"=50*180-{W * (ZP + CD):.1f}-({{N:32-100}}-120)-{{N:32-110}}-{{N:32-120}}-{{N:32-205}}-{{N:32-210}}", calc="lot 9,000 - slab - paving (in lot) - bed - sod", basis="Calculated", conf="Medium")

# ---- 33 Utilities
T("33-100", "Water service 1\" (meter to house)", "", "A100", "Front", "Length", "LF", (0, 0, 14), L=45, basis="Assumed", conf="Low", notes="Meter location unknown (RFI-01).")
T("33-110", "Sewer service 4\" PVC w/ cleanouts", "", "A100", "Front", "Length", "LF", (0, 0, 26), L=50, basis="Assumed", conf="Low")
T("33-120", "Water/sewer tap, meter & impact fees (municipal)", "", "—", "Utility", "Allowance", LS, (0, 0, 3500), count=1, basis="Allowance", conf="Low", notes="Septic system EXCLUDED - assumes public sewer available.")
T("33-130", "Underground electric service lateral", "", "A100", "Site", "Length", "LF", (0, 0, 16), L=60, basis="Assumed", conf="Low")
T("33-140", "Utility company service/connection charges", "", "—", "Utility", "Allowance", LS, (0, 0, 800), count=1, basis="Allowance", conf="Low")

DIVS = {
    "01": "General Requirements", "03": "Concrete", "06": "Wood & Rough Carpentry", "07": "Thermal & Moisture Protection",
    "08": "Openings", "09": "Finishes", "10": "Specialties", "11": "Equipment (Appliances)", "12": "Casework & Countertops",
    "22": "Plumbing", "23": "HVAC", "26": "Electrical", "31": "Earthwork", "32": "Exterior Improvements", "33": "Utilities",
}
DIV_ORDER = ["01", "31", "03", "06", "07", "08", "09", "10", "11", "12", "22", "23", "26", "32", "33"]

# ------------------------------------------------------------------ styling
NAVY = "1F3A5F"
F_HEAD = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
F_INPUT = Font(color="0000CC", name="Calibri", size=10)
F_LINK = Font(color="006100", name="Calibri", size=10)
F_CALC = Font(color="000000", name="Calibri", size=10)
F_BOLD = Font(bold=True, name="Calibri", size=10)
FILL_ASSUMED = PatternFill("solid", fgColor="FFF2CC")
FILL_LOW = PatternFill("solid", fgColor="FCE4D6")
FILL_INPUT = PatternFill("solid", fgColor="DDEBF7")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = F_HEAD; cell.fill = FILL_HEAD; cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def add_table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def page_setup(ws, landscape=True, title_rows=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.center.text = "SFH Development - Preliminary Estimate - Page &P of &N"


# ------------------------------------------------------------------ workbook
wb = Workbook()

# ---------------- Read Me
ws = wb.active; ws.title = "Read Me"
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 100
meta = [
    ("Project", "SFH Development - single-family spec home (Sidney Aulds Building Studio)"),
    ("Owner", "Chris Hill"),
    ("Location", "Ouachita Parish, Louisiana (address TBD - drawings 'Enter address here')"),
    ("Building", "1-story, slab-on-grade, 2 bed / 2 bath. 1,352 gsf conditioned + 117 gsf porch + 504 gsf covered parking = 1,973 gsf (design notes)"),
    ("Estimate purpose / class", "Preliminary budget estimate (AACE Class 4-ish, schematic/progress drawings). NOT a bid. Expected accuracy roughly -15% / +25%."),
    ("Drawings used", "2305031 Progress Drawings dated 5/31/2023 - G000 (cover), A100, A101, A401, A701 - stamped 'NOT FOR CONSTRUCTION'; Spec Home Schematic Design Notes (3 pp)."),
    ("Estimate date / price date", PRICE_DATE),
    ("Units / currency", "US customary units (ft, SF, LF, CY, SQ = 100 SF, EA, LS); USD"),
    ("Pricing basis", "Estimator budget unit costs for the Monroe/Ouachita Parish market, Q4-2026. Material costs are pre-tax. These are NOT supplier or subcontractor quotes - replace with local quotes before contracting."),
    ("Markups", "Sales tax, contingency, overhead and profit are INPUT cells on the Summary sheet (markup on cost). Defaults are placeholders - confirm."),
    ("Status", "PRELIMINARY - quantities and costs require estimator review. Open RFIs listed on 'Assumptions & RFIs'."),
    ("", ""),
    ("Workflow", "Drawing Register -> Takeoff Detail (dims & formulas) -> Cost Database (unit costs w/ source/date) -> Estimate Detail (qty pulled by INDEX/MATCH from takeoff, unit costs from DB) -> Summary (SUMIFS by division + markups) -> Reconciliation (takeoff vs 3D model vs design notes)."),
    ("3D model", "output/model/SFH_3D_Model.html - every model element carries the takeoff ID used in this workbook."),
    ("Regenerate", "python3 scripts/build_estimate.py then scripts/build_all.sh (recalculates with LibreOffice and rebuilds the model)."),
    ("", ""),
    ("COLOR KEY", ""),
    ("Blue text / blue fill", "Hard-coded input (dimension, count, unit price, markup %) - edit these"),
    ("Black text", "Formula / calculation - do not overwrite"),
    ("Green text", "Link pulled from another sheet"),
    ("Yellow row fill", "Basis = Assumed or Allowance (not measured from drawings)"),
    ("Orange row fill", "Confidence = Low - review before use"),
]
ws["A1"] = "SFH DEVELOPMENT - QUANTITY TAKEOFF & PRELIMINARY ESTIMATE"; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
for i, (k, v) in enumerate(meta, 3):
    ws.cell(row=i, column=1, value=k).font = F_BOLD
    c = ws.cell(row=i, column=2, value=v); c.alignment = WRAP
ws["A21"].font = Font(color="0000CC", bold=True); ws["B21"].fill = FILL_INPUT
ws["A23"].fill = FILL_ASSUMED; ws["A24"].fill = FILL_LOW
page_setup(ws, landscape=False)

# ---------------- Drawing Register
ws = wb.create_sheet("Drawing Register")
header(ws, 1, ["Sheet", "Title", "Discipline", "Scale", "Revision / Date", "Status", "File / Page", "Used for", "Notes / Issues"], [9, 26, 13, 14, 18, 18, 30, 40, 60])
reg = [
    ("G000", "Cover sheet (rendering, project team)", "General", "NTS", "Progress 5/31/2023", "Not for construction", "2305031_Progress_Drawings.pdf p1", "Massing check (rendering)", "No sheet number printed on page; design notes call it G000. Project team blank (no structural/MEP engineers)."),
    ("A100", "Site plan", "Architectural", "1/16\" = 1'-0\"", "Progress 5/31/2023", "Not for construction", "p2", "Lot 50' x 180', setbacks, driveway", "Design notes list 'A100 - Plans'; set has A100 = Site Plan, A101 = Plans (numbering conflict, RFI-16). Address 'Enter address here'."),
    ("A101", "Plans: finish floor, finish floor RCP, T.O. plate (roof)", "Architectural", "1/8\" = 1'-0\"", "Progress 5/31/2023", "Not for construction", "p3", "Primary takeoff: overall dims 21'-0\" x 70'-0\" + 24'-0\" carport, room chain, partitions, RCP ceiling heights & fixtures", "Calibrated: 21'-0\" = 395 px @150 dpi (18.81 px/ft); vertical chain 70'-0\" = 1,312 px (18.74 px/ft) - 0.4% check. Interior partitions not dimensioned (SCALED)."),
    ("A401", "Building elevations A-D", "Architectural", "1/4\" = 1'-0\"", "Progress 5/31/2023", "Not for construction", "p4", "T.O. plate 9'-0\", 6/12 pitch, cladding types, window/door tags, gutters", "Overhangs not dimensioned (ASSUMED 12\")."),
    ("A701", "Schedules & details (door/window schedules, types)", "Architectural", "NTS", "Progress 5/31/2023", "Not for construction", "p5", "Door & window counts/sizes", "Door 01 R1 4'-0\" x 8'-0\" leaf/sidelite split not dimensioned."),
    ("Notes", "Spec Home Schematic Design Notes", "Narrative spec", "—", "Undated", "Schematic", "Spec_Home_Schematic_Design_Notes.pdf (3 pp)", "Finishes, framing, sheathing, HVAC, areas", "Areas: porch 117, conditioned 1,352, covered parking 504, total 1,973 gsf."),
    ("—", "Structural drawings", "Structural", "—", "MISSING", "Not issued", "—", "—", "Foundation, beams, trusses: ASSUMED. RFI-05."),
    ("—", "MEP drawings", "MEP", "—", "MISSING", "Not issued", "—", "—", "HVAC, electrical, plumbing priced as design-build allowances."),
]
for r, row in enumerate(reg, 2):
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.alignment = WRAP; cell.border = BORDER
    if row[4] == "MISSING":
        for c in range(1, 10): ws.cell(row=r, column=c).fill = FILL_LOW
ws.freeze_panes = "A2"; page_setup(ws, title_rows="1:1")

# ---------------- Assumptions & RFIs
from assumptions import RFIS  # noqa: E402
ws = wb.create_sheet("Assumptions & RFIs")
header(ws, 1, ["ID", "Category", "Question / Issue", "Assumption used in this estimate", "Qty / cost impact", "Responsible", "Status", "Resolution"], [8, 16, 50, 55, 34, 14, 11, 30])
for r, row in enumerate(RFIS, 2):
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.alignment = WRAP; cell.border = BORDER
    ws.cell(row=r, column=8, value="")
dv = DataValidation(type="list", formula1='"Open,Assumed,Answered,Closed"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"G2:G{len(RFIS) + 1}")
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:H{len(RFIS) + 1}"; page_setup(ws, title_rows="1:1")

# ---------------- Takeoff Detail
ws = wb.create_sheet("Takeoff Detail")
cols = ["ID", "Cost Code", "Division", "Description", "Spec / Reference", "Sheet", "Location", "Takeoff Type", "Count", "Length (ft)",
        "Width / Ht (ft)", "Depth (ft)", "Factor", "Deduct", "Net Qty", "Unit", "Waste %", "Gross Qty", "Round Up?",
        "Calculation / Method", "Basis", "Confidence", "Notes / Assumptions"]
header(ws, 1, cols, [8, 7, 18, 46, 30, 12, 20, 11, 9, 10, 10, 9, 9, 9, 11, 6, 8, 11, 8, 52, 11, 10, 46])
ROWMAP = {}
for i, r in enumerate(ROWS, 2):
    ROWMAP[r["id"]] = i
for i, r in enumerate(ROWS, 2):
    vals = [r["id"], r["code"], DIVS[r["code"]], r["desc"], r["spec"], r["sheet"], r["loc"], r["ttype"],
            r["count"], r["L"], r["Wd"], r["Dp"], r["factor"], r["deduct"]]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font = F_INPUT if 9 <= c <= 14 and v is not None else F_CALC
    # net qty
    if r["net"]:
        f = r["net"]
        for k, rr in ROWMAP.items():
            f = f.replace("{N:" + k + "}", f"O{rr}").replace("{G:" + k + "}", f"R{rr}")
        ws.cell(row=i, column=15, value=f)
    else:
        ws.cell(row=i, column=15, value=f'=IF(I{i}="",1,I{i})*IF(J{i}="",1,J{i})*IF(K{i}="",1,K{i})*IF(L{i}="",1,L{i})*IF(M{i}="",1,M{i})-N(N{i})')
    ws.cell(row=i, column=16, value=r["unit"])
    ws.cell(row=i, column=17, value=r["waste"]).font = F_INPUT
    ws.cell(row=i, column=18, value=f'=IF(S{i}="Y",ROUNDUP(O{i}*(1+Q{i}),0),O{i}*(1+Q{i}))')
    ws.cell(row=i, column=19, value=r["rnd"]).font = F_INPUT
    for c, v in zip(range(20, 24), [r["calc"], r["basis"], r["conf"], r["notes"]]):
        ws.cell(row=i, column=c, value=v)
    for c in range(1, 24):
        cell = ws.cell(row=i, column=c); cell.border = BORDER
        if c in (4, 5, 7, 20, 23): cell.alignment = WRAP
    for c in (9, 10, 11, 12, 13, 14): ws.cell(row=i, column=c).number_format = "#,##0.00"
    ws.cell(row=i, column=15).number_format = "#,##0.00"; ws.cell(row=i, column=18).number_format = "#,##0.00"
    ws.cell(row=i, column=17).number_format = "0%"
NT = len(ROWS) + 1
ws.conditional_formatting.add(f"A2:W{NT}", FormulaRule(formula=[f'$V2="Low"'], fill=FILL_LOW))
ws.conditional_formatting.add(f"A2:W{NT}", FormulaRule(formula=[f'OR($U2="Assumed",$U2="Allowance")'], fill=FILL_ASSUMED))
for rng, lst in ((f"U2:U{NT}", '"Measured,Calculated,Assumed,Allowance,User-provided,External"'), (f"V2:V{NT}", '"High,Medium,Low"'),
                 (f"S2:S{NT}", '"Y,N"'), (f"P2:P{NT}", '"LS,EA,LF,SF,SY,CY,SQ,SHT,RL,MO,WK,SET"')):
    dv = DataValidation(type="list", formula1=lst, allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
add_table(ws, "Takeoff", f"A1:W{NT}")
ws.freeze_panes = "E2"; page_setup(ws, title_rows="1:1")

# ---------------- Cost Database
ws = wb.create_sheet("Cost Database")
header(ws, 1, ["Cost ID", "Cost Code", "Description", "Unit", "Material $/unit", "Labor $/unit", "Sub / Equip / Other $/unit",
               "Total $/unit", "Source", "Market", "Price Date", "Notes"], [8, 7, 46, 6, 11, 11, 13, 11, 50, 26, 11, 30])
for i, r in enumerate(ROWS, 2):
    m, l, s = r["cost"]
    vals = [r["id"], r["code"], r["est_desc"], r["unit"], m, l, s, f"=E{i}+F{i}+G{i}", SRC if not r["id"].startswith("11-100") else "Owner furnished per design notes", MARKET, PRICE_DATE, "Material pre-tax"]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.border = BORDER
        if c in (5, 6, 7): cell.font = F_INPUT; cell.fill = FILL_INPUT; cell.number_format = "$#,##0.00"
        if c == 8: cell.number_format = "$#,##0.00"
        if c in (3, 9): cell.alignment = WRAP
NC = len(ROWS) + 1
add_table(ws, "CostDB", f"A1:L{NC}")
ws.freeze_panes = "D2"; page_setup(ws, title_rows="1:1")

# ---------------- Estimate Detail
ws = wb.create_sheet("Estimate Detail")
header(ws, 1, ["Line", "Cost Code", "Division", "Description", "Takeoff ID", "Qty (gross)", "Unit", "Cost ID", "Material $/u", "Labor $/u",
               "Sub/Equip $/u", "Material $", "Labor $", "Sub/Equip $", "Line Total $", "Basis", "Confidence", "$ / Cond. SF"],
       [6, 7, 20, 46, 9, 11, 6, 8, 11, 11, 11, 12, 12, 12, 13, 11, 10, 10])
TD = "'Takeoff Detail'"; CDB = "'Cost Database'"
for i, r in enumerate(ROWS, 2):
    vals = [i - 1, r["code"], DIVS[r["code"]], r["est_desc"], r["id"],
            f"=INDEX({TD}!$R:$R,MATCH(E{i},{TD}!$A:$A,0))",
            f"=INDEX({TD}!$P:$P,MATCH(E{i},{TD}!$A:$A,0))",
            r["id"],
            f"=INDEX({CDB}!$E:$E,MATCH(H{i},{CDB}!$A:$A,0))",
            f"=INDEX({CDB}!$F:$F,MATCH(H{i},{CDB}!$A:$A,0))",
            f"=INDEX({CDB}!$G:$G,MATCH(H{i},{CDB}!$A:$A,0))",
            f"=F{i}*I{i}", f"=F{i}*J{i}", f"=F{i}*K{i}", f"=L{i}+M{i}+N{i}",
            f"=INDEX({TD}!$U:$U,MATCH(E{i},{TD}!$A:$A,0))",
            f"=INDEX({TD}!$V:$V,MATCH(E{i},{TD}!$A:$A,0))",
            f"=O{i}/Summary!$C$4"]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.border = BORDER
        if c in (6, 7, 9, 10, 11, 16, 17): cell.font = F_LINK
        if c in (9, 10, 11, 12, 13, 14, 15, 18): cell.number_format = "$#,##0.00" if c in (9, 10, 11, 18) else "$#,##0"
        if c == 6: cell.number_format = "#,##0.00"
        if c == 4: cell.alignment = WRAP
NE = len(ROWS) + 1
tr = NE + 1
ws.cell(row=tr, column=4, value="TOTAL DIRECT COST").font = F_BOLD
for c, col in ((12, "L"), (13, "M"), (14, "N"), (15, "O")):
    cell = ws.cell(row=tr, column=c, value=f"=SUBTOTAL(109,{col}2:{col}{NE})"); cell.font = F_BOLD; cell.number_format = "$#,##0"; cell.fill = FILL_TOTAL
ws.conditional_formatting.add(f"A2:R{NE}", FormulaRule(formula=['$Q2="Low"'], fill=FILL_LOW))
ws.conditional_formatting.add(f"A2:R{NE}", FormulaRule(formula=['OR($P2="Assumed",$P2="Allowance")'], fill=FILL_ASSUMED))
add_table(ws, "Estimate", f"A1:R{NE}")
ws.freeze_panes = "E2"; page_setup(ws, title_rows="1:1")

# ---------------- Summary
ws = wb.create_sheet("Summary", 1)
ws.column_dimensions["A"].width = 8; ws.column_dimensions["B"].width = 36
for c in "CDEFGH": ws.column_dimensions[c].width = 15
ws["A1"] = "SFH DEVELOPMENT - PRELIMINARY CONSTRUCTION ESTIMATE"; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws["A2"] = f"Ouachita Parish, LA  |  Progress drawings 5/31/2023 (Not for construction)  |  Price date {PRICE_DATE}  |  Budget estimate - not a bid"
ws["A2"].font = Font(italic=True, size=9)
areas = [("Conditioned area (gsf, design notes)", 1352), ("Front porch (gsf)", 117), ("Covered parking (gsf)", 504)]
ws["B4"] = areas[0][0]; ws["C4"] = 1352
ws["B5"] = areas[1][0]; ws["C5"] = 117
ws["B6"] = areas[2][0]; ws["C6"] = 504
ws["B7"] = "Total under roof (gsf)"; ws["C7"] = "=SUM(C4:C6)"
for r in (4, 5, 6): ws[f"C{r}"].font = F_INPUT; ws[f"C{r}"].fill = FILL_INPUT
ws["E4"] = "Sales tax on materials"; ws["G4"] = 0.10
ws["E5"] = "Contingency (on cost)"; ws["G5"] = 0.05
ws["E6"] = "Overhead (on cost)"; ws["G6"] = 0.08
ws["E7"] = "Profit (on cost)"; ws["G7"] = 0.10
ws["H4"] = "ASSUMED - verify Ouachita rate for address"; ws["H5"] = "placeholder"; ws["H6"] = "placeholder"; ws["H7"] = "placeholder"
for r in (4, 5, 6, 7):
    ws[f"G{r}"].font = F_INPUT; ws[f"G{r}"].fill = FILL_INPUT; ws[f"G{r}"].number_format = "0.0%"; ws[f"H{r}"].font = Font(italic=True, size=8, color="808080")
hr = 9
header_cells = ["Div", "Division", "Material $", "Labor $", "Sub/Equip $", "Total $", "$ / Cond. SF", "% of Direct"]
for i, h in enumerate(header_cells, 1):
    c = ws.cell(row=hr, column=i, value=h); c.font = F_HEAD; c.fill = FILL_HEAD; c.border = BORDER
ED = "'Estimate Detail'"
r0 = hr + 1
for k, d in enumerate(DIV_ORDER):
    r = r0 + k
    ws.cell(row=r, column=1, value=d); ws.cell(row=r, column=2, value=DIVS[d])
    for c, col in ((3, "L"), (4, "M"), (5, "N")):
        ws.cell(row=r, column=c, value=f'=SUMIFS({ED}!${col}$2:${col}${NE},{ED}!$B$2:$B${NE},$A{r})')
    ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})")
    ws.cell(row=r, column=7, value=f"=F{r}/$C$4")
    ws.cell(row=r, column=8, value=f"=F{r}/$F${r0 + len(DIV_ORDER)}")
rl = r0 + len(DIV_ORDER)
ws.cell(row=rl, column=2, value="TOTAL DIRECT COST").font = F_BOLD
for c, col in ((3, "C"), (4, "D"), (5, "E"), (6, "F")):
    ws.cell(row=rl, column=c, value=f"=SUM({col}{r0}:{col}{rl - 1})").font = F_BOLD
ws.cell(row=rl, column=7, value=f"=F{rl}/$C$4"); ws.cell(row=rl, column=8, value=f"=F{rl}/F{rl}")
mk = [
    ("Sales tax on materials", f"=C{rl}*G4"),
    ("Subtotal with tax", f"=F{rl}+F{rl + 1}"),
    ("Contingency", f"=F{rl + 2}*G5"),
    ("Overhead", f"=(F{rl + 2}+F{rl + 3})*G6"),
    ("Profit", f"=(F{rl + 2}+F{rl + 3}+F{rl + 4})*G7"),
    ("GRAND TOTAL (budget)", f"=F{rl + 2}+F{rl + 3}+F{rl + 4}+F{rl + 5}"),
]
for k, (lab, f) in enumerate(mk, 1):
    r = rl + k
    ws.cell(row=r, column=2, value=lab).font = F_BOLD if "TOTAL" in lab or "Subtotal" in lab else F_CALC
    ws.cell(row=r, column=6, value=f)
    ws.cell(row=r, column=7, value=f"=F{r}/$C$4")
GT = rl + len(mk)
ws.cell(row=GT + 1, column=2, value="Grand total per SF under roof (1,973 gsf)")
ws.cell(row=GT + 1, column=6, value=f"=F{GT}/C7")
ws.cell(row=GT + 3, column=2, value="Excluded: land, financing & closing, sales commissions, appliances & W/D (owner), window coverings, fencing, irrigation, septic, gas service, security/AV, solar, architect fees, builder's warranty reserve, escalation beyond price date, rock/unsuitable soil removal, flood elevation above 12\" pad.").alignment = WRAP
ws.merge_cells(start_row=GT + 3, start_column=2, end_row=GT + 3, end_column=8)
ws.row_dimensions[GT + 3].height = 45
for r in range(hr, GT + 2):
    for c in range(1, 9):
        cell = ws.cell(row=r, column=c); cell.border = BORDER
        if c in (3, 4, 5, 6): cell.number_format = "$#,##0"
        if c == 7: cell.number_format = "$#,##0.00"
        if c == 8: cell.number_format = "0.0%"
for r in (rl, GT):
    for c in range(1, 9): ws.cell(row=r, column=c).fill = FILL_TOTAL; ws.cell(row=r, column=c).font = F_BOLD
ws.cell(row=GT, column=6).font = Font(bold=True, size=12)
ws.freeze_panes = "A10"; page_setup(ws, landscape=False)
SUMMARY_GT = f"F{GT}"; SUMMARY_DIRECT = f"F{rl}"

# ---------------- Reconciliation
from reconciliation import recon_rows  # noqa: E402
ws = wb.create_sheet("Reconciliation")
header(ws, 1, ["Check", "Takeoff / Workbook Qty", "Model / Reference Qty", "Unit", "Variance", "Variance %", "Reason / Explanation", "Review Status"],
       [44, 16, 16, 7, 12, 10, 70, 14])
rr = recon_rows(ROWMAP)
for i, (chk, tq, mq, unit, reason, status) in enumerate(rr, 2):
    ws.cell(row=i, column=1, value=chk)
    ws.cell(row=i, column=2, value=tq).font = F_LINK
    ws.cell(row=i, column=3, value=mq).font = F_INPUT
    ws.cell(row=i, column=4, value=unit)
    ws.cell(row=i, column=5, value=f"=B{i}-C{i}")
    ws.cell(row=i, column=6, value=f'=IF(C{i}=0,"",(B{i}-C{i})/C{i})')
    ws.cell(row=i, column=7, value=reason).alignment = WRAP
    ws.cell(row=i, column=8, value=status)
    for c in range(1, 9):
        ws.cell(row=i, column=c).border = BORDER
    for c in (2, 3, 5): ws.cell(row=i, column=c).number_format = "#,##0.00"
    ws.cell(row=i, column=6).number_format = "0.0%"
ws.conditional_formatting.add(f"F2:F{len(rr) + 1}", FormulaRule(formula=[f'AND(ISNUMBER(F2),ABS(F2)>0.05)'], fill=FILL_LOW))
ws.freeze_panes = "B2"; page_setup(ws, title_rows="1:1")

wb.properties.title = "SFH Development - Takeoff & Preliminary Estimate"
wb.properties.creator = "Preliminary estimate (estimator review required)"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", os.path.abspath(OUT), "rows", len(ROWS), "summary GT cell", SUMMARY_GT)
