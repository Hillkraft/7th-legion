"""Build the 3D model HTML and markdown reports from the recalculated workbook.

Usage: python3 scripts/build_reports.py <recalculated.xlsx> <artifact_body_out.html>
"""
import json
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
import model_data as md  # noqa: E402
from assumptions import RFIS  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "output")
xlsx = sys.argv[1]
artifact_out = sys.argv[2] if len(sys.argv) > 2 else None

wb = load_workbook(xlsx, data_only=True)
td = wb["Takeoff Detail"]
TK = {}
for r in range(2, td.max_row + 1):
    i = td.cell(row=r, column=1).value
    if i:
        TK[i] = dict(desc=td.cell(row=r, column=4).value, net=td.cell(row=r, column=15).value or 0, unit=td.cell(row=r, column=16).value,
                     gross=td.cell(row=r, column=18).value or 0, calc=td.cell(row=r, column=20).value, basis=td.cell(row=r, column=21).value,
                     conf=td.cell(row=r, column=22).value, sheet=td.cell(row=r, column=6).value)
ed = wb["Estimate Detail"]
COST = {}
for r in range(2, ed.max_row + 1):
    i = ed.cell(row=r, column=5).value
    if i and i in TK:
        COST[i] = dict(mat=ed.cell(row=r, column=12).value or 0, lab=ed.cell(row=r, column=13).value or 0,
                       sub=ed.cell(row=r, column=14).value or 0, tot=ed.cell(row=r, column=15).value or 0,
                       mu=ed.cell(row=r, column=9).value or 0, lu=ed.cell(row=r, column=10).value or 0, su=ed.cell(row=r, column=11).value or 0)
sm = wb["Summary"]
SUM = []
for r in range(10, 40):
    b = sm.cell(row=r, column=2).value
    f = sm.cell(row=r, column=6).value
    if b and isinstance(f, (int, float)):
        SUM.append((sm.cell(row=r, column=1).value, b, sm.cell(row=r, column=3).value, sm.cell(row=r, column=4).value,
                    sm.cell(row=r, column=5).value, f, sm.cell(row=r, column=7).value))
inputs = {sm[f"E{r}"].value: sm[f"G{r}"].value for r in (4, 5, 6, 7)}
rc = wb["Reconciliation"]
RECON = [[rc.cell(row=r, column=c).value for c in range(1, 9)] for r in range(2, rc.max_row + 1) if rc.cell(row=r, column=1).value]

g = lambda i: TK[i]["gross"]
n = lambda i: TK[i]["net"]
fmt = lambda v, d=0: f"{v:,.{d}f}"
money = lambda v: f"${v:,.0f}"

LAYER_IDS = {
    "site": ["31-140", "32-200", "32-205", "32-207", "32-210", "32-215"],
    "paving": ["32-100", "32-110", "32-120", "32-130"],
    "earth": ["31-100", "31-120", "31-130", "31-150"],
    "slab": ["03-100", "03-130", "03-150", "03-160", "03-170", "03-175", "03-180", "03-185", "03-190", "03-200", "01-180"],
    "beams": ["03-110", "03-120", "03-140", "03-220"],
    "framing_ext": ["06-100", "06-110", "06-120", "06-130"],
    "framing_int": ["06-140", "06-150", "06-160", "06-170", "06-180"],
    "trusses": ["06-190", "06-191", "06-192", "06-230", "06-300"],
    "purlins": ["06-200", "06-210", "06-215", "06-220"],
    "carport": ["06-250", "06-260"],
    "sheathing": ["06-270", "06-290", "06-295"],
    "insulation": ["07-100", "07-110", "07-115"],
    "siding_panel": ["07-150", "07-155"],
    "siding_tg": ["07-160"],
    "ext_ceiling": ["07-170", "07-180"],
    "trim": ["06-240", "06-245", "07-190", "07-200", "07-210", "07-220", "07-225", "07-250", "09-260", "09-265"],
    "roof_sheathing": ["06-280"],
    "shingles": ["07-130", "07-140", "07-145", "07-230", "07-235", "07-240"],
    "windows": ["08-100", "08-110", "08-120", "08-130"],
    "doors": ["08-200", "08-210", "08-220", "08-230", "08-240", "08-250", "08-255", "08-260", "08-265", "08-270", "08-275", "09-270"],
    "drywall": ["09-100", "09-105", "09-110", "09-120", "09-230"],
    "casework": ["12-100", "12-110", "12-120", "12-125", "12-130", "12-200", "12-210", "10-120"],
    "fixtures": ["22-110", "22-120", "22-130", "22-140", "22-145", "22-160", "22-190", "11-110"],
    "mep": ["22-150", "23-100", "23-110", "23-120", "26-100"],
    "lights": ["26-300", "26-310", "26-320", "26-325", "26-330"],
}
QTXT = {
    "site": f"sod {fmt(g('32-210'))} SF · seed {fmt(g('32-215'))} SF · 12 shrubs",
    "paving": f"{fmt(g('32-100') + g('32-110') + g('32-120'))} SF 4\" concrete",
    "earth": f"{fmt(g('31-120'))} CY select fill (loose)",
    "slab": f"{fmt(n('03-160'))} SF · {fmt(g('03-100'), 1)} CY",
    "beams": f"{fmt(g('03-110') + g('03-120'), 1)} CY · {fmt(g('03-140'))} LF #5",
    "framing_ext": f"{fmt(n('06-100'))} LF · {fmt(g('06-120'))} studs 2x6",
    "framing_int": f"{fmt(n('06-140'))} LF · {fmt(g('06-160'))} studs 2x4",
    "trusses": f"{int(g('06-190') + g('06-191') + g('06-192'))} trusses @ 4' O.C.",
    "purlins": f"{fmt(g('06-200'))} LF purlins · {fmt(g('06-210'))} LF furring",
    "carport": f"{fmt(g('06-250'))} LF LVL · {int(g('06-260'))} posts",
    "sheathing": f"{int(g('06-270'))} sheets 4x9 · {int(g('06-290'))} tape rolls",
    "insulation": f"R-21 {fmt(n('07-100'))} SF · R-38 {fmt(n('07-110'))} SF",
    "siding_panel": f"{fmt(n('07-150'))} SF net · {fmt(g('07-150'))} SF gross",
    "siding_tg": f"{fmt(n('07-160'))} SF net · {fmt(g('07-160'))} SF gross",
    "ext_ceiling": f"carport {fmt(n('07-180'))} SF · porch {fmt(n('07-170'))} SF",
    "trim": f"fascia {fmt(n('07-200'))} LF · gutter {fmt(n('07-220'))} LF",
    "roof_sheathing": f"{int(g('06-280'))} sheets 5/8\" ZIP",
    "shingles": f"{fmt(n('07-130'), 1)} SQ net · {fmt(g('07-130'), 1)} SQ gross",
    "windows": "11 EA (5 A · 2 B · 1 C · 3 D)",
    "doors": "3 exterior · 7 interior",
    "drywall": f"walls {fmt(n('09-100'))} SF · ceilings {fmt(n('09-110'))} SF",
    "casework": f"{fmt(n('12-100'), 1)} LF base · {fmt(n('12-110'), 1)} LF wall · {fmt(n('12-200'), 0)} SF tops",
    "fixtures": "2 WC · 2 tubs · 3 sinks · hookups",
    "mep": "2-ton HP · 50 gal WH · 200A",
    "lights": f"{int(n('26-300'))} cans · 2 pendants · 3 fans · 8 ext.",
}
for k, ids in LAYER_IDS.items():
    QTXT[k] += f" · {money(sum(COST[i]['tot'] for i in ids))}"

# ---------------- model HTML
tpl = open(os.path.join(ROOT, "scripts", "model_template.html")).read()
body = tpl.replace("__MODEL_DATA__", md.as_json()).replace("__TAKEOFF_QTY__", json.dumps(QTXT))
os.makedirs(os.path.join(OUT, "model"), exist_ok=True)
# standalone copy inlines three.js r128 (MIT, scripts/vendor) so it opens offline; the artifact copy uses the CDN
vend = os.path.join(ROOT, "scripts", "vendor")
sa_body = body
for url, fn in (("https://cdnjs.cloudflare.com/ajax/libs/three.js/r128/three.min.js", "three.min.js"),
                ("https://cdn.jsdelivr.net/npm/three@0.128.0/examples/js/controls/OrbitControls.js", "OrbitControls.js")):
    if os.path.exists(os.path.join(vend, fn)):
        sa_body = sa_body.replace(f'<script src="{url}"></script>', "<script>\n" + open(os.path.join(vend, fn)).read() + "\n</script>")
standalone = "<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n<meta name=\"viewport\" content=\"width=device-width, initial-scale=1, viewport-fit=cover\">\n" \
    + "<style>:root{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}body{margin:0}[hidden]{display:none!important}</style>\n</head>\n<body>\n" + sa_body + "\n</body>\n</html>\n"
open(os.path.join(OUT, "model", "SFH_3D_Model.html"), "w").write(standalone)
if artifact_out:
    open(artifact_out, "w").write(body)

# ---------------- markdown: assumptions & RFI log
L = ["# Assumptions & RFI Log - SFH Development Spec Home (Ouachita Parish, LA)", "",
     "Source drawings: 2305031 Progress Drawings 5/31/2023 (G000, A100, A101, A401, A701 - NOT FOR CONSTRUCTION) and the Spec Home Schematic Design Notes.",
     "Status key: **Open** = needs an answer before bidding; **Assumed** = estimator assumption carried in the numbers.", "",
     "| ID | Category | Question / issue | Assumption used | Qty / cost impact | Responsible | Status |", "|---|---|---|---|---|---|---|"]
for r in RFIS:
    L.append("| " + " | ".join(str(x).replace("|", "/").replace("\n", " ") for x in r) + " |")
L += ["", "## Building assumptions used for the takeoff", ""]
BA = [
    ("Footprint & dimensions", "21'-0\" wide x 70'-1\" (dimension chain; overall dim reads 70'-0\") + 24'-0\" covered parking; conditioned 21'-0\" x 64'-2\" = 1,347.5 SF (design notes 1,352 gsf); porch 21'-0\" x 5'-11\"."),
    ("Plan scale calibration", "A101 rendered at 150 dpi; 21'-0\" dimension = 395 px (18.81 px/ft); vertical 70'-0\" check = 18.74 px/ft (0.4% agreement). Interior partitions and window centers SCALED (+/- 3\")."),
    ("Wall heights", "T.O. plate 9'-0\" (A401) for all exterior and interior walls; 104-5/8\" precut studs; double top plate; PT bottom plates."),
    ("Ceilings", "Per A101 RCP: 9'-0\" in laundry, hall, master bed, bed #2, kitchen/living; 8'-0\" in storage, master closet, baths, bed #2 closet (dropped 2x4 ceiling joists). Covered parking and porch ceilings slope 6/12 with the roof (vaulted trusses)."),
    ("Framing", "2x6 exterior / 2x4 interior studs @ 24\" O.C.; (2) 2x10 exterior headers; cambered roof trusses @ 48\" O.C. per design notes with 2x4 purlins @ 24\" over and 2x4 ceiling furring @ 24\" under (RFI-06)."),
    ("Roof", "Gable, 6/12, ridge running front-to-back over the full 94'; 12\" eaves and rakes (ASSUMED); ZIP roof (serves as underlayment) + architectural shingles; ridge vent over conditioned attic; vented FC soffit."),
    ("Open gables", "Porch front and carport rear are open gables (end trusses with rake fascia); the gable infill triangles above the laundry wall (faces carport, smooth panel) and living wall (faces porch, T&G) are clad."),
    ("Carport & porch structure", "Posts only at the two carport rear corners; 3-1/2\" x 14\" LVL eave beams (24' clear span) and porch front beam - engineer to size (RFI-07)."),
    ("Foundation", "Monolithic slab-on-grade: 4\" slab, 10 mil poly, #3 @ 18\" EW; 12\"x24\" perimeter beam (4 #5); 12\"x20\" interior beams on ~20' grid; 12\" select-fill pad 5' beyond slab; termite pretreat. No geotech/structural - RFI-04/05."),
    ("Floors", "Finished concrete throughout (design notes): smooth trowel, light grind, densifier, sealer; slab protected during construction. No other flooring."),
    ("Openings", "Counts and sizes per A701 schedules (11 windows, 3 exterior + 7 interior doors), locations scaled from A101 and confirmed on A401 elevations."),
    ("Cladding", "Fiber cement smooth panel on W/E/rear elevations and rear gable infill; fiber cement T&G at porch walls, front gable infill and sloped porch ceiling; FC panel sloped carport ceiling; FC fascia, trim, soffit; primed and field painted."),
    ("Deductions", "Siding and drywall deduct actual opening areas; sheathing, insulation-batt purchase and paint do not add back returns; waste factors are separate per line (5-25%)."),
    ("Insulation", "R-21 batts in 2x6 walls, R-38 blown ceiling (energy code not stated - RFI-03)."),
    ("Interior finishes", "1/2\" GWB walls (MR in baths/laundry), 5/8\" GWB ceilings, Level 3 finish (storage Level 2, unpainted); MDF base/casing/stool & apron painted (not specified - RFI-17)."),
    ("Casework", "Frameless shaker cabinets: 11.7 LF base, 9.4 LF wall, 1 tall pantry, 2 x 30\" vanities; laminate tops 45.6 SF + 4\" splash (scaled - RFI-13)."),
    ("MEP", "Design-build allowances: all-electric, 2-ton heat pump + attic air handler, ducted; 50 gal electric WH; 200A service; 13 plumbing connections; lights counted from RCP (23 cans, 2 pendants, 3 fans) + 8 exterior."),
    ("Site", "Lot 50' x 180' (A100); 12'-wide concrete drive from street (incl. 10' ROW) to carport, 9' x 24' apron at carport, 10.4' x 6' front landing; 4' planting bed with ~12 shrubs; 1 tree; sod front/side, seed rear; public water/sewer."),
    ("Owner items", "Appliances and washer/dryer by owner; hookups by contractor."),
    ("Pricing", f"Estimator budget unit costs, Monroe/Ouachita market, price date 2026-10-15, materials pre-tax. Summary inputs: sales tax {inputs.get('Sales tax on materials', 0):.1%} on materials (verify for address), contingency {inputs.get('Contingency (on cost)', 0):.0%}, overhead {inputs.get('Overhead (on cost)', 0):.0%}, profit {inputs.get('Profit (on cost)', 0):.0%} - all markup on cost, placeholders."),
    ("Schedule", "~20 weeks of general conditions."),
]
for k, v in BA:
    L.append(f"- **{k}:** {v}")
L += ["", "## Exclusions", "",
      "Land; financing, closing and sales costs; appliances & washer/dryer (owner); window coverings; fencing; irrigation; septic system; natural gas service; security/AV; solar; architect fees; builder warranty reserve; escalation after the price date; rock or unsuitable-soil removal; fill above a 12\" average pad; flood-zone elevation requirements."]
open(os.path.join(OUT, "assumptions_rfi_log.md"), "w").write("\n".join(L) + "\n")

# ---------------- markdown: QA & reconciliation
gt = [s for s in SUM if "GRAND" in str(s[1])][0][5]
direct = [s for s in SUM if "DIRECT" in str(s[1])][0]
Q = ["# QA & Reconciliation Report - SFH Development Spec Home", "",
     "Workbook: `output/SFH_Ouachita_Takeoff_Estimate.xlsx` (recalculated with LibreOffice; 0 formula errors).", "",
     "## Takeoff vs. 3D model vs. drawing/notes values", "",
     "| Check | Workbook | Model / reference | Unit | Variance | Var % | Explanation | Status |", "|---|---:|---:|---|---:|---:|---|---|"]
for c in RECON:
    vp = f"{c[5]:.1%}" if isinstance(c[5], (int, float)) else ""
    Q.append(f"| {c[0]} | {c[1]:,.2f} | {c[2]:,.2f} | {c[3]} | {c[4]:,.2f} | {vp} | {c[6] or ''} | {c[7]} |")
Q += ["", "## QA checklist", "",
      "- [x] Every takeoff row has a unit, source sheet, calculation/method, basis and confidence (`Takeoff Detail` columns P, F, T, U, V).",
      "- [x] Plan scale calibrated to a dimensioned reference (21'-0\" on A101) and cross-checked against the 70'-0\" vertical dimension (0.4%).",
      "- [x] Net quantity and waste-adjusted gross quantity are kept in separate columns; waste is a visible input per line.",
      "- [x] Opening deductions applied to siding and drywall; not to sheathing (cut waste) - documented per row.",
      "- [x] Window/door counts reconcile: A701 schedule = A101 tags = A401 elevations = 3D model (11 windows, 10 doors).",
      "- [x] Every estimate line pulls quantity from the takeoff by ID (INDEX/MATCH) and unit cost from the Cost Database; allowances are labeled as such.",
      "- [x] Each cost has a source note, market and price date; all are estimator budget values, not quotes.",
      "- [x] Summary totals by division via SUMIFS reconcile to the Estimate Detail subtotal "
      f"({money(direct[5])} direct).",
      "- [x] Markups are explicit inputs (markup on cost) and are not buried in unit prices.",
      "- [ ] Human estimator review of scaled partition positions and kitchen layout (RFI-13).",
      "- [ ] Replace budget unit costs with local supplier/sub quotes.",
      "- [ ] Structural, MEP and geotechnical information (RFI-04/05/06/07/14/15).", "",
      "## Low-confidence items", ""]
for i, t in TK.items():
    if t["conf"] == "Low":
        Q.append(f"- `{i}` {t['desc']} - {fmt(t['gross'], 2)} {t['unit']} ({money(COST[i]['tot'])})")
Q += ["", "## Notes on known differences", "",
      "- Interior wall drywall is taken to 9'-0\" in all rooms; rooms with 8'-0\" ceilings are over-measured by about 1' of wall height (~130 SF) - conservative.",
      "- The model represents studs, trusses and purlins at their nominal spacing for visualization; the stud count in the workbook adds corners, intersections and cripples, so the model shows fewer studs than are purchased.",
      "- The design-notes porch area (117 gsf) is a clear area; the takeoff uses the slab area (124.25 SF) for concrete and finishing.",
      f"- Grand total {money(gt)} includes placeholder markups; with owner acting as GC (no overhead/profit) the total is {money(direct[5] + [s for s in SUM if 'tax' in str(s[1]).lower()][0][5] + [s for s in SUM if 'Contingency' in str(s[1])][0][5])}."]
open(os.path.join(OUT, "qa_reconciliation.md"), "w").write("\n".join(Q) + "\n")

# ---------------- markdown: estimate summary (for README / quick read)
S = ["# Preliminary Estimate Summary - SFH Development Spec Home", "",
     "Ouachita Parish, Louisiana · progress drawings 5/31/2023 (not for construction) · price date 2026-10-15 · budget estimate, not a bid.", "",
     "| Div | Division | Material | Labor | Sub/Equip | Total | $/cond. SF |", "|---|---|---:|---:|---:|---:|---:|"]
for s in SUM:
    if s[2] is not None:
        S.append(f"| {s[0] or ''} | {s[1]} | {money(s[2])} | {money(s[3])} | {money(s[4])} | **{money(s[5])}** | ${s[6]:,.2f} |")
    else:
        S.append(f"| | {s[1]} | | | | **{money(s[5])}** | {'$' + format(s[6], ',.2f') if isinstance(s[6], (int, float)) else ''} |")
open(os.path.join(OUT, "estimate_summary.md"), "w").write("\n".join(S) + "\n")

# ---------------- model notes
MN = ["# 3D Model Notes", "",
      "- **File:** `output/model/SFH_3D_Model.html` (open in any modern browser; three.js r128 is inlined, so it works offline).",
      "- **Application:** custom three.js (WebGL) model generated by `scripts/build_reports.py` from `scripts/model_data.py` + `scripts/model_template.html`. No Revit/SketchUp/Blender was available in this environment, and none was requested.",
      "- **Units / axes:** feet. x = west->east (0 = west face of framing), z = rear->front (0 = laundry wall face; carport -24..0; porch 64.17..70.08), y = up (0 = finish floor).",
      "- **Layers (toggle individually, by group, or by view preset):** site, paving, select-fill pad, slab, grade beams, exterior & interior framing, trusses, purlins/furring, carport beams/posts, ZIP sheathing, insulation, panel siding, T&G siding, exterior ceilings, trim/soffit/gutters, roof sheathing, shingles, windows, doors, drywall, casework, plumbing fixtures/appliances, MEP equipment, light fixtures, room labels.",
      "- **Interaction:** drag to orbit, scroll/pinch to zoom, right-drag to pan; click any element for its takeoff ID, spec, size and basis; hide an element or isolate its layer; horizontal section-cut slider; plan view.",
      "- **Traceability:** every element carries the workbook takeoff ID (e.g. `06-120`), and each layer row shows the takeoff quantity and the estimate dollars for that layer.",
      "- **Limitations:** visual coordination model only, not BIM. Structure (foundation beams, LVLs, headers, truss webs) is schematic and ASSUMED (no structural drawings). Studs are drawn at nominal 24\" spacing with kings/jacks/headers/cripples. Truss web layout is representative. Partition and fixture locations are scaled from A101. MEP routing (duct trunk) is assumed. Overhangs assumed 12\".",
      "- **Verified dimensions (model vs. plan):** width 21'-0\"; laundry 6'-10\"; master 13'-5\"; baths 9'-10\"; bed #2 11'-2\"; kitchen/living 22'-11\"; porch 5'-11\"; carport 24'-0\"; plate 9'-0\"; pitch 6/12 - all set from dimensioned values in `model_data.py`."]
open(os.path.join(OUT, "model", "model_notes.md"), "w").write("\n".join(MN) + "\n")
print("grand total", money(gt), "direct", money(direct[5]))
